import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
p = doc.add_paragraph("Test Paragraph")
table = doc.add_table(rows=2, cols=2)
cell = table.cell(0, 0)
cell.text = "Header"
shading = parse_xml(r'<w:shd {} w:fill="1A365D"/>'.format(nsdecls('w')))
cell._tc.get_or_add_tcPr().append(shading)
doc.save("test_output.docx")
print("Saved test_output.docx successfully!")
