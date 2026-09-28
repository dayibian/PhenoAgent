import os
import time
import docx
from docx.shared import Inches, Pt, RGBColor

t0 = time.time()
grid = "R290040058"
ehr_path = f"data/stuttering/ehr_markdown_dataset/{grid}.md"

with open(ehr_path, "r", encoding="utf-8") as f:
    text = f.read()

doc = docx.Document()
doc.add_heading(f"Patient {grid} - Test", level=1)

# Add paragraphs
lines = text.split("\n")
current_chunk = []
for line in lines:
    if line.startswith("### "):
        if current_chunk:
            doc.add_paragraph("\n".join(current_chunk))
            current_chunk = []
        h = doc.add_heading(line.replace("### ", ""), level=2)
    else:
        current_chunk.append(line)
if current_chunk:
    doc.add_paragraph("\n".join(current_chunk))

out_test = "test_large_patient.docx"
doc.save(out_test)
print(f"Created {out_test} in {time.time() - t0:.2f}s, size: {os.path.getsize(out_test)} bytes")
