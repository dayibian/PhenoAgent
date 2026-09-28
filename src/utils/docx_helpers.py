import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border_left_accent(cell, color_hex="2B6CB0", sz="36"):
    # sz="36" is 4.5pt thick border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(borders)

def set_table_borders(table, color="D2D6DC", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def set_box_borders(cell, color="CBD5E0", sz="8"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tcBorders>')
    tcPr.append(borders)

def add_callout_box(doc, text_list, title="Extracted EHR Evidence & Clinical Mentions", bg_hex="F0F4F8", accent_hex="2B6CB0"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_border_left_accent(cell, accent_hex, sz="24")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(26, 54, 93)
    
    for item in text_list:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.left_indent = Inches(0.15)
        run_bullet = p2.add_run("• ")
        run_bullet.bold = True
        run_bullet.font.name = "Calibri"
        run_bullet.font.size = Pt(9.5)
        run_bullet.font.color.rgb = RGBColor(43, 108, 176)
        
        run_text = p2.add_run(item)
        run_text.font.name = "Calibri"
        run_text.font.size = Pt(9.5)
        run_text.font.color.rgb = RGBColor(45, 55, 72)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_expert_adjudication_form(doc, patient_grid, current_dx):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "FAF5FF" if current_dx == "Excluded" else "F7FAFC")
    set_box_borders(cell, color="CBD5E0", sz="10")
    set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_h = p.add_run(f"📋 EXPERT ADJUDICATION & REVIEW SIGN-OFF — Patient {patient_grid}")
    run_h.bold = True
    run_h.font.name = "Calibri"
    run_h.font.size = Pt(10)
    run_h.font.color.rgb = RGBColor(44, 82, 130)
    
    options = [
        f"☐ AGREE with PhenoAgent: Classify as {current_dx}",
        "☐ OVERTURN to POSITIVE: Patient exhibits true developmental / persistent stuttering",
        "☐ OVERTURN to INDETERMINATE: Conflicting evidence or insufficient clinical detail",
        "☐ OTHER: Please specify in reviewer notes below"
    ]
    for opt in options:
        p_opt = cell.add_paragraph()
        p_opt.paragraph_format.space_before = Pt(1)
        p_opt.paragraph_format.space_after = Pt(2)
        p_opt.paragraph_format.left_indent = Inches(0.15)
        r = p_opt.add_run(opt)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(45, 55, 72)
        
    p_notes = cell.add_paragraph()
    p_notes.paragraph_format.space_before = Pt(6)
    p_notes.paragraph_format.space_after = Pt(2)
    r_n = p_notes.add_run("Expert Reviewer Notes & Clinical Rationale:")
    r_n.bold = True
    r_n.font.name = "Calibri"
    r_n.font.size = Pt(9.5)
    
    # 2 blank lines for notes
    for _ in range(2):
        p_line = cell.add_paragraph()
        p_line.paragraph_format.space_before = Pt(0)
        p_line.paragraph_format.space_after = Pt(2)
        r_l = p_line.add_run("_________________________________________________________________________________")
        r_l.font.color.rgb = RGBColor(203, 213, 224)
        r_l.font.size = Pt(9)
        
    # Signature line
    p_sig = cell.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(4)
    p_sig.paragraph_format.space_after = Pt(0)
    r_sig = p_sig.add_run("Reviewer Signature / Initials: ____________________     Date: ____________     Credentials: _________")
    r_sig.font.name = "Calibri"
    r_sig.font.size = Pt(9)
    r_sig.font.color.rgb = RGBColor(74, 85, 104)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

print("Helper functions ready.")
