import os
import re
import json
import glob
import yaml
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Paths
results_csv = "/home/biand/Projects/PhenoAgent/results/controls_151/controls_151_results.csv"
json_dir = "/home/biand/Projects/PhenoAgent/results/controls_151/json_results"
gt_path = "/home/biand/Projects/Celiac_BioVU/data/Celiac Diagnosis by Manual Review.xlsx"
controls_path = "/home/biand/Projects/Celiac_BioVU/data/controls_need_review.csv"
dataset_dir = "/home/biand/Projects/Celiac_BioVU/data/ehr_markdown_dataset"
yaml_path = "/home/biand/Projects/Celiac_BioVU/data/celiac_keywords_latest.yaml"
out_dir = "/home/biand/Projects/PhenoAgent/data/additional_notes_for_review"

os.makedirs(out_dir, exist_ok=True)

# 1. Load results and ground truth
df_res = pd.read_csv(results_csv)
pos = df_res[df_res['diagnosis'] == 'Positive'].copy()

wb = openpyxl.load_workbook(gt_path, data_only=True)
sheet = wb.active
gt_grids = set()
for row in sheet.iter_rows(min_row=2, values_only=True):
    g = str(row[0]).strip() if row[0] is not None else ''
    if g:
        gt_grids.add(g)

# Filter to 33 grids without ground truth
unreviewed_pos = pos[~pos['grid'].isin(gt_grids)].copy()
unreviewed_pos = unreviewed_pos.sort_values(by='grid').reset_index(drop=True)

# Merge with ICD counts from controls_need_review.csv
df_ctrl = pd.read_csv(controls_path)
unreviewed_pos = pd.merge(unreviewed_pos, df_ctrl[['grid', 'icd_count']], on='grid', how='left')
unreviewed_pos['icd_count'] = unreviewed_pos['icd_count'].fillna(0).astype(int)

print(f"Targeting {len(unreviewed_pos)} unreviewed positive grids.")

# Load keywords for filtering notes
with open(yaml_path, 'r', encoding='utf-8') as f:
    kw_data = yaml.safe_load(f)
popular_words = kw_data.get('popular_words', [])
kw_pattern = re.compile('|'.join([re.escape(w) for w in popular_words]), re.IGNORECASE) if popular_words else None

def parse_patient_markdown(md_path):
    if not os.path.exists(md_path):
        return {'labs': [], 'notes': []}
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Labs
    labs = []
    labs_match = re.search(r'## Labs\n(.*?)(## Medical Notes|\Z)', content, re.DOTALL)
    if labs_match:
        labs_text = labs_match.group(1).strip()
        for line in labs_text.split('\n'):
            line = line.strip()
            if line.startswith('-'):
                labs.append(line.replace('- ', '').strip())
            elif line and line != 'No labs available.':
                labs.append(line)
                
    # Medical Notes
    notes = []
    notes_match = re.search(r'## Medical Notes\n(.*)', content, re.DOTALL)
    if notes_match:
        notes_text = notes_match.group(1).strip()
        splits = re.split(r'###\s+\[', notes_text)
        for sp in splits:
            sp = sp.strip()
            if not sp:
                continue
            parts = sp.split(']', 1)
            if len(parts) < 2:
                continue
            date_str = parts[0].strip()
            rest = parts[1].strip()
            lines = rest.split('\n')
            title = lines[0].strip()
            source = "Unknown"
            body_start = 1
            if len(lines) > 1 and lines[1].strip().startswith('**Source:**'):
                source = lines[1].replace('**Source:**', '').strip()
                body_start = 2
            body = "\n".join(lines[body_start:]).strip()
            notes.append({
                'date': date_str,
                'title': title,
                'source': source,
                'body': body
            })
            
    filtered_notes = []
    if kw_pattern:
        for n in notes:
            text_to_check = f"{n['title']} {n['source']} {n['body']}"
            if kw_pattern.search(text_to_check):
                filtered_notes.append(n)
    if not filtered_notes and notes:
        filtered_notes = notes
        
    return {'labs': labs, 'notes': filtered_notes, 'total_notes_count': len(notes)}

# 2. Build individual markdown review files and single master markdown report
master_md = "# Clinical Review Package: Unreviewed Positive Cases (N=33)\n\n"
master_md += "This document contains patient profiles, automated agent diagnostic logic, extracted evidence, and relevant medical notes for the **33 patients** predicted as **Positive** by PhenoAgent who currently lack manual ground truth review in `Celiac Diagnosis by Manual Review.xlsx`.\n\n"
master_md += "## Summary Table of Cases\n\n"
master_md += "| # | Patient ID | ICD Count | Agent Prediction | Confidence | Extracted Evidence Summary |\n"
master_md += "| - | ---------- | --------- | ---------------- | ---------- | -------------------------- |\n"

excel_rows = []

for idx, row in unreviewed_pos.iterrows():
    grid = row['grid']
    conf = row['confidence']
    reasoning = str(row['reasoning'])
    evidence = str(row['evidence']) if pd.notna(row['evidence']) else ""
    icd_cnt = row['icd_count']
    
    # Load JSON for extra details if available
    jf_path = os.path.join(json_dir, f"{grid}.json")
    signals = []
    if os.path.exists(jf_path):
        with open(jf_path) as f:
            jd = json.load(f)
            # Find positive signals
            if 'trace' in jd and 'steps' in jd['trace']:
                for st in jd['trace']['steps']:
                    if st.get('agent') == 'SignalExtractor' and 'signals' in st:
                        for sig in st['signals']:
                            if sig.get('past_celiac_diagnosis') or sig.get('external_confirmation') or sig.get('iel_status') == 'positive' or sig.get('marsh_grade') == 'positive':
                                signals.append(sig)

    md_path = os.path.join(dataset_dir, f"{grid}.md")
    clinical_data = parse_patient_markdown(md_path)
    
    ev_summary = evidence.replace("\n", " ")[:120] + ("..." if len(evidence) > 120 else "")
    master_md += f"| {idx+1} | [{grid}](#patient-{grid.lower()}) | {icd_cnt} | Positive | {conf:.2f} | {ev_summary} |\n"
    
    # Excel record
    excel_rows.append({
        "Patient ID": grid,
        "ICD Count": icd_cnt,
        "Agent Prediction": "Positive",
        "Confidence": conf,
        "Key Extracted Evidence": evidence,
        "Agent Reasoning Chain": reasoning,
        "Relevant Notes Count": len(clinical_data['notes']),
        "Total Notes Count": clinical_data.get('total_notes_count', len(clinical_data['notes'])),
        "Clinician Diagnosis (Yes/No/PMH/Unknown)": "",
        "Clinician Comments / Rationale": ""
    })
    
    # Write Individual Patient Review File
    indiv_md = f"# Clinical Review Packet: Patient {grid}\n\n"
    indiv_md += f"- **Patient Identifier:** `{grid}`\n"
    indiv_md += f"- **Cohort:** Controls Need Review (`no_ttgiga`)\n"
    indiv_md += f"- **Historical Celiac ICD Billing Count:** {icd_cnt}\n"
    indiv_md += f"- **Agent Prediction:** **Positive** (Confidence: `{conf:.2f}`)\n"
    indiv_md += f"- **Ground Truth Status:** Unreviewed (Not in manual review spreadsheet)\n\n"
    
    indiv_md += "## Clinical Adjudication Sign-Off\n\n"
    indiv_md += "- [ ] **Positive (Yes)** — Confirmed Celiac Disease (Biopsy proven or serology + clinical confirmation)\n"
    indiv_md += "- [ ] **Negative (No)** — Not Celiac Disease (Competing diagnosis, normal biopsies, misdiagnosis, screening only)\n"
    indiv_md += "- [ ] **PMH Only** — Historical mention in problem list without objective confirmation\n"
    indiv_md += "- [ ] **Unknown / Indeterminate** — Insufficient clinical documentation\n\n"
    indiv_md += "**Clinician Notes & Comments:**  \n"
    indiv_md += "> \n\n"
    
    indiv_md += "---\n\n"
    indiv_md += "## Agent Diagnostic Reasoning\n\n"
    indiv_md += f"**Decision Logic & Reasoning:**  \n{reasoning}\n\n"
    if evidence:
        indiv_md += f"**Extracted Quotes & Evidence:**  \n> {evidence}\n\n"
        
    indiv_md += "## Laboratory Results\n\n"
    if clinical_data['labs']:
        for l in clinical_data['labs']:
            indiv_md += f"- {l}\n"
    else:
        indiv_md += "*No relevant laboratory values recorded in electronic health record.*\n"
    indiv_md += "\n"
    
    indiv_md += f"## Relevant Medical Notes ({len(clinical_data['notes'])} notes containing celiac-relevant keywords)\n\n"
    for n_idx, n in enumerate(clinical_data['notes'], 1):
        indiv_md += f"### Note {n_idx}: [{n['date']}] {n['title']}\n"
        indiv_md += f"**Source:** *{n['source']}*\n\n"
        indiv_md += f"```text\n{n['body']}\n```\n\n"
        indiv_md += "---\n\n"
        
    indiv_file_path = os.path.join(out_dir, f"{grid}_review.md")
    with open(indiv_file_path, 'w', encoding='utf-8') as f:
        f.write(indiv_md)

master_md += "\n---\n\n"
master_md += "## Detailed Patient Profiles and Reasoning\n\n"

for idx, row in unreviewed_pos.iterrows():
    grid = row['grid']
    conf = row['confidence']
    reasoning = str(row['reasoning'])
    evidence = str(row['evidence']) if pd.notna(row['evidence']) else ""
    icd_cnt = row['icd_count']
    
    md_path = os.path.join(dataset_dir, f"{grid}.md")
    clinical_data = parse_patient_markdown(md_path)
    
    master_md += f"### Patient {grid}\n\n"
    master_md += f"- **Patient Identifier:** `{grid}`\n"
    master_md += f"- **Agent Prediction:** **Positive** | **Confidence:** `{conf:.2f}` | **ICD Count:** {icd_cnt}\n"
    master_md += f"- **Individual Review File:** [{grid}_review.md](./{grid}_review.md)\n\n"
    master_md += f"**Extracted Evidence:**  \n> {evidence}\n\n"
    master_md += f"**Agent Reasoning:**  \n{reasoning}\n\n"
    
    master_md += "#### Relevant Medical Notes Excerpts\n\n"
    for n_idx, n in enumerate(clinical_data['notes'][:5], 1):  # top 5 notes in master report
        master_md += f"**Note {n_idx}: [{n['date']}] {n['title']}** (*{n['source']}*)\n"
        excerpt = n['body'][:500] + ("..." if len(n['body']) > 500 else "")
        master_md += f"> {excerpt.replace(chr(10), ' ')}\n\n"
    if len(clinical_data['notes']) > 5:
        master_md += f"*... and {len(clinical_data['notes']) - 5} additional notes in [{grid}_review.md](./{grid}_review.md).*\n\n"
    master_md += "---\n\n"

master_file_path = os.path.join(out_dir, "clinical_review_summary.md")
with open(master_file_path, 'w', encoding='utf-8') as f:
    f.write(master_md)
print(f"Generated master markdown report at: {master_file_path}")

# 3. Create formatted Excel scoring sheet for clinicians
df_excel = pd.DataFrame(excel_rows)
csv_sheet_path = os.path.join(out_dir, "clinical_review_sheet.csv")
df_excel.to_csv(csv_sheet_path, index=False)

excel_sheet_path = os.path.join(out_dir, "clinical_review_sheet.xlsx")
wb_out = openpyxl.Workbook()
ws = wb_out.active
ws.title = "Clinical Review (33 Patients)"

# Header formatting
header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

cols = list(df_excel.columns)
ws.append(cols)

for col_idx, col_name in enumerate(cols, 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Highlight review columns in light yellow for clinicians
review_fill = PatternFill(start_color="FEFCBF", end_color="FEFCBF", fill_type="solid")

for row_idx, r_data in enumerate(excel_rows, 2):
    for col_idx, col_name in enumerate(cols, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=r_data[col_name])
        cell.font = Font(name="Calibri", size=10)
        cell.border = border_thin
        if "Clinician" in col_name:
            cell.fill = review_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        else:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

# Adjust column widths
col_widths = {
    "A": 16, # Patient ID
    "B": 12, # ICD Count
    "C": 16, # Agent Prediction
    "D": 12, # Confidence
    "E": 35, # Key Extracted Evidence
    "F": 50, # Agent Reasoning Chain
    "G": 18, # Relevant Notes Count
    "H": 16, # Total Notes Count
    "I": 28, # Clinician Diagnosis
    "J": 45  # Clinician Comments
}
for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

wb_out.save(excel_sheet_path)
print(f"Generated clinician review Excel sheet at: {excel_sheet_path}")
print(f"Generated 33 individual review files in: {out_dir}")
